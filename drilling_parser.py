"""Read the inputs a Belloy Well Report (.xlsm) holds on its `Data` and `DSum1`/`DSum2` sheets."""
from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field

import openpyxl

MAX_LEGS = 10
DAILY_FIRST_ROW, DAILY_LAST_ROW = 8, 71          # DSum1 / DSum2 daily rows


class WellReportError(ValueError):
    """The uploaded file is not a Well Report in the expected layout."""


@dataclass
class Leg:
    number: int
    kop: float | None          # start (kick-off) MD
    td: float | None           # end MD
    reached: dt.datetime | None

    @property
    def length(self) -> float:
        if self.kop is None or self.td is None:
            return 0.0
        return max(0.0, self.td - self.kop)


@dataclass
class DailyRow:
    date: dt.datetime
    time: dt.time | None
    from_md: float | None
    to_md: float | None
    operations: str
    drilled: float | None = None          # the report's own DRILLED column (L)
    day: float | None = None              # DAY column (D)


@dataclass
class WellInputs:
    well: dict
    spud: dt.datetime | None
    surface_casing_md: float | None
    surface_casing_reached: dt.datetime | None
    heel_md: float | None
    heel_reached: dt.datetime | None
    td_md: float | None
    td_reached: dt.datetime | None
    legs: list[Leg]
    daily: list[DailyRow]
    warnings: list[str] = field(default_factory=list)
    drill_out_surface: dt.datetime | None = None      # Data row 27
    int_casing_md: float | None = None                # Data row 31
    int_casing_reached: dt.datetime | None = None
    legs_declared: int | None = None                  # Data D39 'Total # of Legs'


# ---------------------------------------------------------------- helpers
def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(',', '').strip())
        except ValueError:
            return None
    return None


def _date(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime.combine(v, dt.time())
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
    return None


def _time(v):
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    if isinstance(v, dt.timedelta):
        return (dt.datetime.min + v).time()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v < 1:
        return (dt.datetime.min + dt.timedelta(days=float(v))).time().replace(microsecond=0)
    if isinstance(v, str) and ':' in v:
        try:
            h, m = v.strip().split(':')[:2]
            return dt.time(int(h) % 24, int(m))
        except ValueError:
            return None
    return None


def _stamp(d, t):
    d, t = _date(d), _time(t)
    if d is None or t is None:
        return None
    return dt.datetime.combine(d.date(), t)


def _text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _grid(ws, max_row, max_col):
    """Read a sheet block into {(row, col): value} (1-based)."""
    out = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True), 1):
        for c, v in enumerate(row, 1):
            if v is not None:
                out[(r, c)] = v
    return out


def idle_days(report_dates):
    """Calendar days with no daily report between the first and last report (e.g. the well was suspended)."""
    have = {d.date() if isinstance(d, dt.datetime) else d for d in report_dates}
    if not have:
        return set()
    first, last = min(have), max(have)
    return {first + dt.timedelta(days=i) for i in range((last - first).days + 1)} - have


def active_hours(t0, t1, idle):
    """Hours from t0 to t1, leaving out whole days on which no report was written."""
    if t0 is None or t1 is None:
        return None
    skip = sum(1 for d in idle if t0.date() < d < t1.date())
    return (t1 - t0).total_seconds() / 3600 - 24 * skip


# ---------------------------------------------------------------- main entry
def read_well_report(file_bytes: bytes) -> WellInputs:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise WellReportError(f'Could not open the file as an Excel workbook ({exc}).') from exc
    try:
        if 'Data' not in wb.sheetnames or 'DSum1' not in wb.sheetnames:
            raise WellReportError('This workbook has no "Data" and "DSum1" sheets, so it is not a Well Report.')
        D = _grid(wb['Data'], 80, 22)
        warnings: list[str] = []

        # find cells by their labels, so older / newer Well Report layouts (rows shifted) still read
        def label_rows(col, prefix):
            p = prefix.lower()
            return [r for (r, c), v in sorted(D.items()) if c == col and isinstance(v, str)
                    and v.strip().lower().startswith(p)]

        def row_of(col, *prefixes):
            for pr in prefixes:
                rows = label_rows(col, pr)
                if rows:
                    return rows[0]
            return None

        def val(row, col):
            return None if row is None else D.get((row, col))

        B, C, Dc, E, G, I, K = 2, 3, 4, 5, 7, 9, 11
        r_name = row_of(B, 'Well Name')
        leg_rows = label_rows(B, 'Hz Leg')
        if r_name is None or not leg_rows:
            raise WellReportError('The "Data" sheet does not look like a Well Report: no "Well Name:" label or '
                                  '"Hz Leg ... Interval" rows were found in column B.')
        r_elev = [r for (r, c), v in D.items() if c == I and isinstance(v, str) and v.strip().lower().startswith('ground')]
        r_kb = [r for (r, c), v in D.items() if c == I and isinstance(v, str) and v.strip().lower().startswith('kelly')]
        well = {
            'Licensee': _text(val(row_of(B, 'Licensee'), C)), 'Well Name': _text(val(r_name, C)),
            'UWI': _text(val(row_of(B, 'U.W.I', 'UWI'), C)),
            'Licence #': _text(val(row_of(B, 'Licence #', 'License #', 'Licence No', 'License No'), C)),
            'AFE #': _text(val(row_of(B, 'A.F.E', 'AFE'), C)),
            'Province': _text(val(row_of(G, 'Province'), I)), 'Field': _text(val(row_of(G, 'Field'), I)),
            'Target Zone': _text(val(row_of(G, 'Target Zone'), I)),
            'Drilling Contractor': _text(val(row_of(B, 'Drilling Contractor'), C)),
            'Surface Location': _text(val(row_of(G, 'Surface Location'), I)),
            'Bottom Hole Location': _text(val(row_of(G, 'Bottom Hole Location'), I)),
            'Ground Elevation (m)': _num(val(r_elev[0] if r_elev else None, K)),
            'KB Elevation (m)': _num(val(r_kb[0] if r_kb else None, K)),
        }

        def stamp(row):
            return _stamp(val(row, I), val(row, K))

        legs = []
        for n, r in enumerate(leg_rows[:MAX_LEGS], 1):          # numbered by position (labels are sometimes stale)
            legs.append(Leg(n, _num(val(r, C)), _num(val(r, E)), stamp(r)))
        legs = [lg for lg in legs if lg.kop or lg.td]

        r_spud = row_of(G, 'SPUD')
        spud = stamp(r_spud)
        if spud is None and val(r_spud, I) not in (None, '', '-'):
            warnings.append(f'Spud date is not a date ("{_text(val(r_spud, I))}"), so Surface time is left blank.')
        r_surf = row_of(B, 'Surface Casing Depth')
        r_int = row_of(B, 'Intermediate Casing')
        r_heel = row_of(B, 'Hz Heel Point')
        heel_md, heel_t = _num(val(r_heel, C)), stamp(r_heel)
        int_md, int_t = _num(val(r_int, C)), stamp(r_int)
        if heel_md is None and int_md is not None:
            heel_md, heel_t = int_md, int_t
            warnings.append(f'Hz Heel Point is blank on "Data"; the intermediate casing point ({int_md:,.0f} m) '
                            'is used as the heel.')
        r_td = row_of(B, 'Total Depth')
        td_md, td_t = _num(val(r_td, C)), stamp(r_td)
        if td_md is None and legs:
            last = max((lg for lg in legs if lg.td), key=lambda lg: lg.reached or dt.datetime.min, default=None)
            if last is not None:
                td_md, td_t = last.td, last.reached
                warnings.append(f'No "Total Depth" row on "Data"; Leg {last.number} TD ({td_md:,.0f} m) is used as the well TD.')
        r_legs = row_of(B, 'Total # of Legs')
        declared = _num(val(r_legs, Dc)) or _num(val(r_legs, C))

        daily: list[DailyRow] = []
        for sheet in ('DSum1', 'DSum2'):
            if sheet not in wb.sheetnames:
                continue
            S = _grid(wb[sheet], DAILY_LAST_ROW, 15)
            for r in range(DAILY_FIRST_ROW, DAILY_LAST_ROW + 1):
                d, to = _date(S.get((r, 1))), _num(S.get((r, 10)))
                if d is None and to is None:
                    continue
                if d is None:
                    warnings.append(f'{sheet} row {r}: depth {to:g} m has no date and was skipped.')
                    continue
                daily.append(DailyRow(d, _time(S.get((r, 6))), _num(S.get((r, 8))), to, _text(S.get((r, 15))),
                                      _num(S.get((r, 12))), _num(S.get((r, 4)))))
        if not daily:
            raise WellReportError('No daily progress rows were found on "DSum1" (date in column A, TO depth in column J).')
        daily.sort(key=lambda x: x.date)
        days = [x.day for x in daily if x.day is not None]
        dupes = sorted({int(v) for v in days if days.count(v) > 1})
        if dupes:
            warnings.append('DSum1 DAY column repeats day ' + ', '.join(map(str, dupes))
                            + '; the day number here is counted from the dates instead.')

        return WellInputs(
            well=well, spud=spud,
            surface_casing_md=_num(val(r_surf, C)), surface_casing_reached=stamp(r_surf),
            heel_md=heel_md, heel_reached=heel_t, td_md=td_md, td_reached=td_t,
            legs=legs, daily=daily, warnings=warnings,
            drill_out_surface=stamp(row_of(G, 'Drill Out Surf')),
            int_casing_md=int_md, int_casing_reached=int_t,
            legs_declared=int(declared) if declared else None,
        )
    finally:
        wb.close()

